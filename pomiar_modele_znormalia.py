# -*- coding: utf-8 -*-
"""pomiar_modele_znormalia.py — mapowanie znormaliów na modele 3D z pary OUT + IAM.

Po co: znormalia (łożyska, siłowniki, paski) nie mają numeru rysunku ani IDW,
więc nie da się dla nich odczytać modelu tak jak dla detali RMPAK (te mają
numer rysunku wprost w iProperty „Part Number").

Skąd wtedy wziąć model: normalia siedzą w złożeniach. Zestawiamy dwa źródła
tego samego złożenia:

    OUT.xlsx / ELEMENTY ZNORMALIZOWANE  — nazwa wg tabelki IDW + ilość
    IAM przez ApprenticeServer          — plik modelu + liczba wystąpień

Nazwa („688ZZ") powstaje w tabelce IDW, nie w modelu, więc w IAM-ie nie ma
jej wprost — stąd dopasowanie po kodzie katalogowym i ilości, nie po kluczu.

⚠️ Kolumna „Pliki 3D" w OUT to NIE ścieżka, tylko flaga „STP" (istnieje
eksport STEP). Nie da się jej użyć.

Wynik pomiaru na B:\\!BIBLIOTEKA (27.09.2026): 66% pozycji dopasowanych
pewnie, 34% zostaje człowiekowi — nazwa w tabelce nie ma tam nic wspólnego
z nazwą modelu („SP-1,2/10,6/30 Sprężyna…" → „sprężyna klawiszy wewn.ipt").

CZYSTY ODCZYT — nic nie zapisuje do baz. ApprenticeServer to osobny, lekki
proces: nie rusza sesji Inventora ani dokumentów otwartych przez użytkownika.

Użycie:
    python pomiar_modele_znormalia.py --root "B:\\!BIBLIOTEKA"
    python pomiar_modele_znormalia.py --root "B:\\" --json pary.json
"""
from __future__ import annotations

import argparse
import json
import os
import re
import sys
import time
from collections import Counter, defaultdict

# ZZ / 2RS / 2Z to typ uszczelnienia łożyska, nie część kodu. Bez obcięcia
# „6004ZZ" nie trafia w „6004 2Z Łożysko….ipt" — ta jedna poprawka podniosła
# skuteczność pomiaru z 52% na 66%.
SUFIKS_LOZYSKA = re.compile(r"(2rs1?|2z|rs1?|zz|z|n|k|c3)$", re.IGNORECASE)

PROG_PRZYJECIA = 12   # minimum punktów, żeby uznać dopasowanie
PROG_PRZEWAGI = 6     # o tyle musi wygrać z drugim kandydatem


def zbitka(t: str) -> str:
    """Tylko litery i cyfry, małe. „6004 ZZ" -> „6004zz"."""
    return re.sub(r"[^0-9a-ząćęłńóśźż]+", "", (t or "").lower())


def rdzen(kod: str) -> str:
    """„6004zz" -> „6004". Obcinamy tylko gdy zostaje sam numer łożyska."""
    bez = SUFIKS_LOZYSKA.sub("", kod)
    return bez if re.fullmatch(r"\d{3,6}", bez) else kod


def kody(t: str) -> set:
    """Fragmenty wyglądające na kod katalogowy: min 4 znaki, zawierają cyfrę."""
    out = set()
    for tok in re.split(r"[\s,;/()]+", (t or "").lower()):
        tok = zbitka(tok.strip(".-_"))
        if len(tok) >= 4 and re.search(r"\d", tok):
            out.add(tok)
            out.add(rdzen(tok))
    return {k for k in out if len(k) >= 3 and re.search(r"\d", k)}


def slowa(t: str) -> set:
    return set(re.findall(r"[a-ząćęłńóśźż]{4,}", (t or "").lower()))


def znormalia_z_out(path):
    """[(nazwa, ilość całkowita)] z arkusza ELEMENTY ZNORMALIZOWANE."""
    import openpyxl
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception:
        return []
    if "ELEMENTY ZNORMALIZOWANE" not in wb.sheetnames:
        wb.close()
        return []
    ws = wb["ELEMENTY ZNORMALIZOWANE"]
    pozycje, kol = [], {}
    # Arkusz ma sekcję na każde podzłożenie, więc nagłówek powtarza się
    # wielokrotnie — kolumny czytamy za każdym razem od nowa.
    for row in ws.iter_rows(values_only=True):
        vals = ["" if c is None else str(c).strip() for c in row[:12]]
        if "Poz." in vals and "Nazwa" in vals:
            kol = {v: i for i, v in enumerate(vals) if v}
            continue
        if not kol:
            continue
        try:
            nazwa = vals[kol["Nazwa"]]
            ilosc = vals[kol.get("Ilość całkowita", kol.get("Ilość"))]
        except (KeyError, IndexError):
            continue
        if not nazwa or nazwa == "Nazwa":
            continue
        try:
            q = int(round(float(ilosc.replace(",", "."))))
        except (ValueError, AttributeError):
            continue
        pozycje.append((nazwa, q))
    wb.close()
    return pozycje


def komponenty_iam(app, iam):
    """{ścieżka: (liczba wystąpień, nazwa, part number)} dla CAŁEGO drzewa."""
    try:
        dok = app.Open(iam)
    except Exception:
        return None
    licznik, nazwy, numery = Counter(), {}, {}

    def zejdz(occ):
        try:
            d = occ.Definition.Document
            s = d.FullFileName
        except Exception:
            return
        pn = ""
        try:
            pn = d.PropertySets.Item("Design Tracking Properties").Item("Part Number").Value or ""
        except Exception:
            pass
        licznik[s] += 1
        nazwy.setdefault(s, occ.Name)
        numery.setdefault(s, pn.strip())
        try:
            for sub in occ.SubOccurrences:
                zejdz(sub)
        except Exception:
            pass

    try:
        for occ in dok.ComponentDefinition.Occurrences:
            zejdz(occ)
    finally:
        try:
            app.Close()
        except Exception:
            pass
    return {s: (licznik[s], nazwy[s], numery[s]) for s in licznik}


def ocen(nazwa_out, ilosc_out, nazwa_komp, plik, ilosc_komp):
    """Punkty dopasowania + czytelny powód. Kod katalogowy waży najwięcej."""
    tekst = nazwa_komp + " " + os.path.splitext(os.path.basename(plik))[0]
    zb_out, zb_komp = zbitka(nazwa_out), zbitka(tekst)
    pkt, powod = 0, []

    if len(zb_out) >= 5 and zb_out in zb_komp:
        pkt += 20
        powod.append("symbol w nazwie")
    else:
        wspolne = kody(nazwa_out) & kody(tekst)
        if wspolne:
            najdluzszy = max(wspolne, key=len)
            pkt += 12 + 2 * (len(wspolne) - 1) + min(len(najdluzszy), 8)
            powod.append("kod " + najdluzszy)
        else:
            czesciowe = [k for k in kody(nazwa_out) if len(k) >= 4 and k in zb_komp]
            if czesciowe:
                pkt += 9
                powod.append("kod~" + czesciowe[0])

    wspolne_slowa = slowa(nazwa_out) & slowa(tekst)
    if wspolne_slowa:
        pkt += 2 * len(wspolne_slowa)
        powod.append("słowa")
    if ilosc_komp == ilosc_out:
        pkt += 3
        powod.append("ilość")
    return pkt, "; ".join(powod)


def main():
    ap = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--root", default=r"B:\!BIBLIOTEKA", help="katalog do przeskanowania")
    ap.add_argument("--json", help="zapisz pary symbol -> model do pliku JSON")
    args = ap.parse_args()

    sys.stdout.reconfigure(encoding="utf-8")
    import win32com.client as w

    outy, iamy = [], {}
    for dirpath, _d, files in os.walk(args.root):
        if "oldversion" in dirpath.lower():
            continue
        for f in files:
            low = f.lower()
            if low.endswith("_out.xlsx"):
                outy.append(os.path.join(dirpath, f))
            elif low.endswith(".iam"):
                iamy.setdefault(os.path.splitext(f)[0].strip().lower(),
                                os.path.join(dirpath, f))
    print(f"katalog: {args.root}")
    print(f"OUT-ów: {len(outy)}   IAM-ów: {len(iamy)}\n")

    app = w.Dispatch("Inventor.ApprenticeServer")
    t0 = time.time()
    pewne, niepewne = [], []
    pary = defaultdict(set)
    bez_iam = 0

    for out in sorted(outy):
        baza = os.path.basename(out)[:-len("_OUT.xlsx")]
        # nazwa IAM = nazwa OUT bez numeru rysunku z przodu
        trzon = re.sub(r"^[A-Za-z0-9]{1,8}-\d{3}\.\d{2,3}[A-Za-z]*\s+", "",
                       baza).strip().lower()
        iam = iamy.get(trzon) or iamy.get(baza.lower())
        pozycje = znormalia_z_out(out)
        if not pozycje:
            continue
        if not iam:
            bez_iam += 1
            continue
        komponenty = komponenty_iam(app, iam)
        if komponenty is None:
            continue
        # Kandydaci: komponenty BEZ Part Number. Detale RMPAK mają tam numer
        # rysunku, więc odpadają — i dobrze, bo one mapują się wprost.
        kandydaci = [(s, q, nm) for s, (q, nm, pn) in komponenty.items() if not pn]

        for nazwa, ilosc in pozycje:
            oceny = []
            for s, q, nm in kandydaci:
                pkt, powod = ocen(nazwa, ilosc, nm, s, q)
                if pkt:
                    oceny.append((pkt, s, nm, powod))
            oceny.sort(key=lambda x: (-x[0], x[1]))
            if not oceny:
                niepewne.append((nazwa, ilosc, "brak sygnału"))
                continue
            najlepszy = oceny[0]
            drugi = oceny[1][0] if len(oceny) > 1 else 0
            if najlepszy[0] >= PROG_PRZYJECIA and najlepszy[0] >= drugi + PROG_PRZEWAGI:
                pewne.append((nazwa, ilosc, najlepszy[1], najlepszy[3]))
                pary[nazwa].add(najlepszy[1])
            else:
                niepewne.append((nazwa, ilosc, f"{najlepszy[2][:24]} ({najlepszy[0]} pkt)"))

    razem = len(pewne) + len(niepewne)
    if not razem:
        print("Nie znaleziono pozycji znormalizowanych.")
        return
    konflikty = {n: sorted(v) for n, v in pary.items() if len(v) > 1}

    print(f"czas: {time.time() - t0:.0f}s   złożeń bez IAM: {bez_iam}")
    print("=" * 66)
    print(f"POZYCJI ZNORMALIZOWANYCH: {razem}")
    print(f"  pewne     {len(pewne):>5}  ({len(pewne) / razem * 100:.0f}%)   automat")
    print(f"  niepewne  {len(niepewne):>5}  ({len(niepewne) / razem * 100:.0f}%)   człowiek")
    print(f"\n  unikalnych symboli zmapowanych: {len(pary)}")
    print(f"  z tego z konfliktem (2+ modele): {len(konflikty)}")
    print("=" * 66)

    print("\n--- ZMAPOWANE ---")
    for n in sorted(pary):
        modele = sorted(pary[n])
        znak = "   !! KONFLIKT" if len(modele) > 1 else ""
        print(f"   {n[:34]:36} -> {os.path.basename(modele[0])[:36]}{znak}")

    print("\n--- DO RĘCZNEGO PRZYPISANIA ---")
    for n in sorted({x[0] for x in niepewne}):
        print(f"   {n[:60]}")

    if args.json:
        with open(args.json, "w", encoding="utf-8") as f:
            json.dump({n: sorted(v) for n, v in pary.items()}, f,
                      ensure_ascii=False, indent=1)
        print(f"\npary zapisane: {args.json}")


if __name__ == "__main__":
    main()
